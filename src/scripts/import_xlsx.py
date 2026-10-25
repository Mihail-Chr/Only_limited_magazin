"""
Парсит xlsx Ozon в product + listing.
Запуск:
  python -m src.scripts.import_xlsx
  python -m src.scripts.import_xlsx --file "scratch/ozon_reports/Видеоигра_24.10.2026.xlsx"
"""
import argparse
from pathlib import Path
import openpyxl
from src.config import settings, get_store
from src.db import get_conn


STORE_CODE = "limited_games"
WAREHOUSE_CODE = "main_warehouse"
MAX_HEADER_SCAN = 20  # сколько первых строк просматривать в поисках шапки


def pick(row: dict, key: str) -> str | None:
    v = row.get(key)
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def to_int(v: str | None) -> int | None:
    if not v:
        return None
    try:
        return int(float(str(v).replace(",", ".")))
    except ValueError:
        return None


def to_float(v: str | None) -> float | None:
    if not v:
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def find_header(ws) -> tuple[int, list[str]]:
    """Ищет строку, где есть 'Артикул' (с * или без). Возвращает (row_idx, headers)."""
    for r in range(1, MAX_HEADER_SCAN + 1):
        values = [c.value for c in ws[r]]
        headers = [str(v).strip() if v is not None else "" for v in values]
        if any(h.startswith("Артикул") for h in headers):
            return r, headers
    return 0, []


def find_sheet(wb):
    """Приоритет: лист 'Шаблон'. Иначе — первый лист, где есть 'Артикул' в первых 20 строках."""
    if "Шаблон" in wb.sheetnames:
        ws = wb["Шаблон"]
        header_row, headers = find_header(ws)
        if header_row:
            return ws, header_row, headers
    for name in wb.sheetnames:
        ws = wb[name]
        header_row, headers = find_header(ws)
        if header_row:
            return ws, header_row, headers
    # не нашли — вернём первый лист с диагностикой
    ws = wb[wb.sheetnames[0]]
    return ws, 0, [c.value for c in ws[1]]


def diagnose(ws, header_row: int, headers: list[str]) -> None:
    print("\n--- ДИАГНОСТИКА ---")
    print(f"Лист: {ws.title}")
    print(f"Всего строк: {ws.max_row}, колонок: {ws.max_column}")
    print(f"Найденная строка шапки: {header_row or 'НЕ НАЙДЕНА'}")
    print("Заголовки первых 30 колонок:")
    for i, h in enumerate(headers[:30], start=1):
        print(f"  {i:>2}. {h!r}")
    print("\nПервые 3 непустые строки после шапки (или с начала):")
    start = (header_row + 1) if header_row else 1
    shown = 0
    for r in range(start, ws.max_row + 1):
        row = [c.value for c in ws[r]]
        if any(v is not None for v in row):
            print(f"  строка {r}: {row[:15]}")
            shown += 1
            if shown >= 3:
                break
    print("--- /ДИАГНОСТИКА ---\n")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", help="Путь к xlsx. Если не указан — ищется в scratch/ozon_reports")
    args = ap.parse_args()

    if args.file:
        xlsx = Path(args.file)
        if not xlsx.is_absolute():
            xlsx = (settings.data_raw_xlsx.parent.parent / xlsx).resolve()
    else:
        candidates = sorted(settings.data_raw_xlsx.glob("*.xlsx"))
        if not candidates:
            raise SystemExit(f"Нет xlsx в {settings.data_raw_xlsx}")
        # предпочитаем файл, в имени которого есть 'игр'
        preferred = [p for p in candidates if "игр" in p.name.lower()]
        xlsx = (preferred or candidates)[0]

    if not xlsx.exists():
        raise SystemExit(f"Файл не найден: {xlsx}")
    print(f"Парсим: {xlsx}")

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws, header_row, headers = find_sheet(wb)

    if not header_row:
        print("Шапка не найдена — показываю диагностику, чтобы понять структуру.")
        diagnose(ws, header_row, headers)
        raise SystemExit(
            "Не нашёл строку с 'Артикул' в первых "
            f"{MAX_HEADER_SCAN} строках. Смотрите диагностику выше."
        )

    idx = {h: i for i, h in enumerate(headers) if h}
    print(f"Лист: {ws.title}, шапка: строка {header_row}, колонок: {len(idx)}")

    store = get_store(STORE_CODE)
    with get_conn() as conn:
        store_row = conn.execute(
            "SELECT id FROM store WHERE code=?", (store.code,)
        ).fetchone()
        if not store_row:
            raise SystemExit("Сначала выполните seed.py")
        store_id = store_row["id"]

        wh_row = conn.execute(
            "SELECT id FROM warehouse WHERE code=?", (WAREHOUSE_CODE,)
        ).fetchone()
        if not wh_row:
            raise SystemExit("Склад не найден, выполните seed.py")
        wh_id = wh_row["id"]

        inserted = 0
        skipped = 0
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not any(v is not None for v in row):
                continue
            d = {h: row[i] for h, i in idx.items() if i < len(row)}
            offer_id = pick(d, "Артикул*") or pick(d, "Артикул")
            name = pick(d, "Название товара")
            if not offer_id or not name:
                skipped += 1
                continue

            conn.execute(
                """INSERT INTO product(
                    offer_id, name, brand, product_type, platform, edition,
                    age_rating, genre, publisher, developer, media_type,
                    media_count, release_year, language, view, modes,
                    weight_g, width_mm, height_mm, length_mm,
                    photo_main_url, photo_extra_urls, country_of_origin,
                    package_type, package_size
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(offer_id) DO UPDATE SET
                    name=excluded.name,
                    brand=excluded.brand,
                    product_type=excluded.product_type,
                    platform=excluded.platform,
                    edition=excluded.edition,
                    age_rating=excluded.age_rating,
                    genre=excluded.genre,
                    publisher=excluded.publisher,
                    developer=excluded.developer,
                    media_type=excluded.media_type,
                    media_count=excluded.media_count,
                    release_year=excluded.release_year,
                    language=excluded.language,
                    view=excluded.view,
                    modes=excluded.modes,
                    weight_g=excluded.weight_g,
                    width_mm=excluded.width_mm,
                    height_mm=excluded.height_mm,
                    length_mm=excluded.length_mm,
                    photo_main_url=excluded.photo_main_url,
                    photo_extra_urls=excluded.photo_extra_urls,
                    country_of_origin=excluded.country_of_origin,
                    package_type=excluded.package_type,
                    package_size=excluded.package_size,
                    updated_at=datetime('now')""",
                (
                    offer_id, name,
                    pick(d, "Бренд*"),
                    pick(d, "Тип*"),
                    pick(d, "Платформа*"),
                    pick(d, "Издание"),
                    pick(d, "Возрастные ограничения*"),
                    pick(d, "Жанр игры"),
                    pick(d, "Издатель"),
                    pick(d, "Разработчик"),
                    pick(d, "Тип носителя"),
                    to_int(pick(d, "Количество носителей, шт.")),
                    to_int(pick(d, "Год выпуска")),
                    pick(d, "Язык"),
                    pick(d, "Вид"),
                    pick(d, "Режимы"),
                    to_int(pick(d, "Вес в упаковке, г*")),
                    to_int(pick(d, "Ширина упаковки, мм*")),
                    to_int(pick(d, "Высота упаковки, мм*")),
                    to_int(pick(d, "Длина упаковки, мм*")),
                    pick(d, "Ссылка на главное фото*"),
                    pick(d, "Ссылки на дополнительные фото"),
                    pick(d, "Страна-изготовитель"),
                    pick(d, "Упаковка"),
                    pick(d, "Размер упаковки (Длина х Ширина х Высота), см"),
                ),
            )
            product_id = conn.execute(
                "SELECT id FROM product WHERE offer_id=?", (offer_id,)
            ).fetchone()["id"]

            conn.execute(
                """INSERT INTO listing(
                    store_id, product_id, warehouse_id, sku, barcode,
                    price, old_price, vat, status, error_text, defects_text
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(store_id, product_id) DO UPDATE SET
                    price=excluded.price,
                    old_price=excluded.old_price,
                    vat=excluded.vat,
                    error_text=excluded.error_text,
                    defects_text=excluded.defects_text""",
                (
                    store_id, product_id, wh_id,
                    pick(d, "SKU"),
                    pick(d, "Штрихкод (Серийный номер / EAN)"),
                    to_float(pick(d, "Предельная цена, руб.*")),
                    to_float(pick(d, "Зачёркнутая цена, руб.")),
                    pick(d, "НДС, %*"),
                    "draft",
                    pick(d, "Ошибка"),
                    pick(d, "Недочёты"),
                ),
            )

            conn.execute(
                """INSERT INTO stock(warehouse_id, product_id, quantity)
                   VALUES(?, ?, 0)
                   ON CONFLICT(warehouse_id, product_id) DO NOTHING""",
                (wh_id, product_id),
            )
            inserted += 1

        conn.commit()
        print(f"Обработано строк: {inserted}")
        if skipped:
            print(f"Пропущено (нет артикула или названия): {skipped}")


if __name__ == "__main__":
    main()